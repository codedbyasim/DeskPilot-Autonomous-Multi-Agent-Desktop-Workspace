"""
DeskPilot — PDF Reading & Extraction Tools
Uses pdfplumber for resilient text, table, and invoice extraction.
"""

from strands import tool
import pdfplumber
import json
import re
from pathlib import Path
from datetime import datetime
from typing import Optional, Union, Dict, Any, List
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from backend.config.settings import DEFAULT_OUTPUT_DIR
from backend.utils.document_resolver import resolve_document_path
from backend.utils.logger import get_logger

logger = get_logger("tools.pdf")


def _resolve_pdf_path(file_path: str) -> Path:
    """
    Resolves relative or friendly PDF paths across cwd, sample_data, Desktop, Documents, and output directories.
    """
    if not file_path or not str(file_path).strip():
        return Path("sample_data")

    resolved = resolve_document_path(file_path, default_ext=".pdf")
    if resolved:
        return resolved

    p_str = str(file_path).strip().strip("'\"")
    if p_str.startswith("~"):
        return Path(p_str).expanduser().resolve()

    p_norm = p_str.replace("\\", "/").strip()
    if p_norm.lower() == "desktop":
        return (Path.home() / "Desktop").resolve()
    if p_norm.lower().startswith("desktop/"):
        return (Path.home() / "Desktop" / p_norm[8:]).resolve()

    p = Path(p_str)
    if p.exists():
        return p.resolve()

    # Fallback to desktop or output
    desktop_cand = Path.home() / "Desktop" / p.name
    if desktop_cand.exists():
        return desktop_cand.resolve()

    output_cand = DEFAULT_OUTPUT_DIR / p.name
    if output_cand.exists():
        return output_cand.resolve()

    return p.resolve()


@tool
def list_pdfs_in_folder(folder_path: str = "") -> str:
    """
    List all PDF files in a given folder with file paths and sizes.

    Args:
        folder_path: Path to folder to scan. Defaults to 'sample_data' or 'Desktop'.

    Returns:
        Formatted list of PDF file paths found
    """
    target_str = folder_path.strip() if folder_path else "sample_data"
    folder = _resolve_pdf_path(target_str)
    logger.info(f"Scanning for PDFs in: {folder}")

    if not folder.exists():
        return f"Folder not found: '{folder_path}' (resolved: '{folder}')"

    if not folder.is_dir():
        return f"Path is not a directory: '{folder_path}' (resolved: '{folder}')"

    pdfs = sorted(list(folder.glob("*.pdf")), key=lambda x: x.name.lower())
    if not pdfs:
        return f"No PDF files found in '{folder}'"

    result = f"PDF files in '{folder}':\n{'─'*60}\n"
    for i, p in enumerate(pdfs, 1):
        size_kb = p.stat().st_size // 1024
        try:
            rel_path = p.relative_to(Path.cwd()).as_posix()
        except ValueError:
            rel_path = p.as_posix()
        result += f"  {i}. {rel_path} ({size_kb} KB)\n"

    logger.info(f"Found {len(pdfs)} PDFs in {folder}")
    result += f"\nTotal: {len(pdfs)} PDF file(s)"
    return result


@tool
def read_pdf(file_path: str) -> str:
    """
    Read and extract all text from a PDF file page by page.

    Args:
        file_path: Absolute or relative path to the PDF file

    Returns:
        Extracted text content from the PDF
    """
    path = _resolve_pdf_path(file_path)
    logger.info(f"Reading PDF: {path}")

    if not path.exists():
        return f"Error: PDF file not found at '{file_path}' (resolved: '{path}')"
    if not path.suffix.lower() == ".pdf":
        return f"Error: File is not a PDF: '{file_path}' (extension: '{path.suffix}')"

    try:
        pages_content = []
        total_text_chars = 0
        with pdfplumber.open(path) as pdf:
            total_pages = len(pdf.pages)
            for i, page in enumerate(pdf.pages, 1):
                page_text = page.extract_text() or ""
                clean_text = page_text.strip()
                total_text_chars += len(clean_text)
                pages_content.append(f"--- Page {i}/{total_pages} ---\n{clean_text if clean_text else '[No digital text on this page]'}")

        if total_text_chars == 0:
            return f"PDF Content: {path.name}\n{'='*60}\nNotice: No digital text found across {total_pages} page(s). This document may be a scanned image."

        full_text = "\n\n".join(pages_content)
        logger.info(f"Extracted {total_text_chars} chars from {path.name} ({total_pages} pages)")
        return f"PDF Content: {path.name} ({total_pages} pages)\n{'='*60}\n{full_text}"

    except Exception as e:
        logger.error(f"PDF read failed: {e}")
        return f"Error reading PDF '{path.name}': {str(e)}"


@tool
def extract_pdf_tables(file_path: str) -> str:
    """
    Extract all tables from a PDF file as structured JSON and Markdown preview.

    Args:
        file_path: Path to the PDF file

    Returns:
        JSON and Markdown tables extracted from the PDF
    """
    path = _resolve_pdf_path(file_path)
    logger.info(f"Extracting tables from PDF: {path}")

    if not path.exists():
        return f"Error: PDF file not found at '{file_path}' (resolved: '{path}')"

    try:
        all_tables = []
        md_tables = []
        with pdfplumber.open(path) as pdf:
            for page_idx, page in enumerate(pdf.pages, 1):
                tables = page.extract_tables()
                if tables:
                    for t_idx, raw_table in enumerate(tables):
                        cleaned_rows = []
                        for row in raw_table:
                            if not row or all(c is None or str(c).strip() == "" for c in row):
                                continue
                            cleaned_row = [str(c).strip().replace("\n", " ") if c is not None else "" for c in row]
                            cleaned_rows.append(cleaned_row)

                        if cleaned_rows:
                            all_tables.append({
                                "page": page_idx,
                                "table_index": t_idx,
                                "rows": cleaned_rows
                            })

                            if len(cleaned_rows) >= 1:
                                headers = cleaned_rows[0]
                                md = f"\nTable (Page {page_idx}, #{t_idx+1}):\n"
                                md += "| " + " | ".join(headers) + " |\n"
                                md += "| " + " | ".join(["---"] * len(headers)) + " |\n"
                                for r in cleaned_rows[1:]:
                                    padded = r + [""] * (len(headers) - len(r))
                                    md += "| " + " | ".join(padded[:len(headers)]) + " |\n"
                                md_tables.append(md)

        if not all_tables:
            return f"No tables found in '{path.name}'."

        result_json = json.dumps(all_tables, indent=2, ensure_ascii=False)
        preview = "\n".join(md_tables)
        logger.info(f"Found {len(all_tables)} tables in {path.name}")
        return f"Extracted Tables from {path.name}:\n{preview}\n\nStructured JSON:\n{result_json}"

    except Exception as e:
        logger.error(f"Table extraction failed: {e}")
        return f"Error extracting tables from '{path.name}': {str(e)}"


@tool
def extract_pdf_invoice(file_path: str) -> str:
    """
    Extract structured invoice data from a PDF — vendor, invoice number, date, total, line items.
    Returns structured JSON for reconciliation and financial reporting workflows.

    Args:
        file_path: Path to invoice PDF

    Returns:
        JSON with invoice fields: vendor, invoice_number, date, total, line_items, tables
    """
    path = _resolve_pdf_path(file_path)
    logger.info(f"Extracting invoice data from: {path}")

    if not path.exists():
        return f"Error: File not found: '{file_path}' (resolved: '{path}')"

    try:
        text_lines = []
        raw_tables = []

        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text() or ""
                text_lines.extend([line.strip() for line in page_text.splitlines() if line.strip()])
                tables = page.extract_tables()
                if tables:
                    raw_tables.extend(tables)

        full_text = "\n".join(text_lines)

        v_match = (
            re.search(r"(?:Vendor|Supplier|From|Company|Billed By):\s*([^\n\r]+)", full_text, re.IGNORECASE)
            or re.search(r"TAX INVOICE\s*\n+([^\n\r]+)", full_text, re.IGNORECASE)
        )
        inv_match = re.search(r"(?:Invoice\s*(?:Number|No\.?|#|ID)?|Bill\s*No\.?):\s*([^\n\r]+)", full_text, re.IGNORECASE)
        d_match = re.search(r"(?:Invoice\s*Date|Date|Dated):\s*([^\n\r]+)", full_text, re.IGNORECASE)
        t_match = (
            re.search(r"(?:Total\s*Amount\s*Due|Grand\s*Total|Total\s*Due|Amount\s*Due|Total):\s*\$?([0-9,]+\.?[0-9]*)", full_text, re.IGNORECASE)
            or re.search(r"\$\s*([0-9,]+\.[0-9]{2})", full_text)
        )

        vendor = v_match.group(1).strip() if v_match else "Unknown Vendor"
        invoice_number = inv_match.group(1).strip() if inv_match else path.stem
        invoice_date = d_match.group(1).strip() if d_match else ""
        total_amount = float(t_match.group(1).replace(",", "")) if t_match else 0.0

        line_items = []
        cleaned_tables = []
        for tbl in raw_tables:
            if not tbl or len(tbl) < 2:
                continue
            cleaned_tbl = []
            for row in tbl:
                cleaned_row = [str(c).strip().replace("\n", " ") if c is not None else "" for c in row]
                cleaned_tbl.append(cleaned_row)
            cleaned_tables.append(cleaned_tbl)

            for row in tbl[1:]:
                if not row or all(c is None or str(c).strip() == "" for c in row):
                    continue
                cells = [str(c).strip() if c else "" for c in row]
                if len(cells) >= 4:
                    line_items.append({
                        "description": cells[0],
                        "quantity": cells[1],
                        "unit_price": cells[2],
                        "total": cells[3]
                    })
                elif len(cells) >= 2:
                    line_items.append({
                        "description": cells[0],
                        "amount": cells[-1]
                    })

        invoice_data = {
            "file": path.name,
            "file_path": path.as_posix(),
            "vendor": vendor,
            "invoice_number": invoice_number,
            "date": invoice_date,
            "total": total_amount,
            "line_items": line_items,
            "tables": cleaned_tables,
            "line_count": len(text_lines),
        }

        logger.info(f"Invoice extracted: {vendor}, {invoice_number}, ${total_amount:,.2f} ({len(line_items)} line items)")
        return json.dumps(invoice_data, indent=2, ensure_ascii=False)

    except Exception as e:
        logger.error(f"Invoice extraction failed: {e}")
        return f"Error extracting invoice from '{path.name}': {str(e)}"


@tool
def edit_pdf_document(
    file_path: str,
    instructions: str = "",
    replacements: Optional[Union[dict, str]] = None,
    append_content: str = "",
    output_path: str = "",
) -> str:
    """
    Read, analyze, and edit content from an existing PDF document.
    Extracts text, headings, and data tables using pdfplumber, applies your requested
    text replacements and updates, and generates an updated, professionally formatted
    deliverable document (Word .docx format for instant viewing and editing).
    User can provide just a filename (e.g. 'TB_Lab_Report.pdf' or 'invoice_001.pdf').

    Args:
        file_path: Path or filename of the PDF to edit
        instructions: Description of modifications performed
        replacements: Dict or JSON string of {'old text': 'new text'} to replace in the document
        append_content: Markdown content to append (headings #, bullet points, and tables | Col 1 | Col 2 |)
        output_path: Destination path for the edited document (e.g. 'Desktop/Updated_Report.docx')

    Returns:
        Status message with deliverable location, changes summary, and verification report.
    """
    from backend.tools.word_tools import (
        _parse_input_to_sections, _apply_table_styling,
        _resolve_word_path, verify_word_document
    )

    path = _resolve_pdf_path(file_path)
    logger.info(f"Editing PDF document: {path}")

    if not path.exists():
        return f"Error: PDF file not found at '{file_path}' (resolved: '{path}')"
    if path.suffix.lower() != ".pdf":
        return f"Error: File '{path.name}' is not a PDF document."

    # Parse replacements
    replace_map = {}
    if replacements:
        if isinstance(replacements, str):
            try:
                parsed = json.loads(replacements)
                if isinstance(parsed, dict):
                    replace_map = {str(k): str(v) for k, v in parsed.items() if str(k)}
            except Exception:
                pass
        elif isinstance(replacements, dict):
            replace_map = {str(k): str(v) for k, v in replacements.items() if str(k)}

    # Extract text and tables from PDF
    extracted_sections = []
    extracted_tables = []
    replacement_count = 0

    try:
        with pdfplumber.open(path) as pdf:
            for page_idx, page in enumerate(pdf.pages, 1):
                raw_text = page.extract_text() or ""
                if raw_text.strip():
                    mod_text = raw_text.strip()
                    for old_txt, new_txt in replace_map.items():
                        if old_txt in mod_text:
                            count = mod_text.count(old_txt)
                            mod_text = mod_text.replace(old_txt, new_txt)
                            replacement_count += count
                    extracted_sections.append({
                        "page": page_idx,
                        "content": mod_text
                    })

                page_tables = page.extract_tables()
                if page_tables:
                    for t_idx, raw_table in enumerate(page_tables):
                        cleaned_rows = []
                        for row in raw_table:
                            if not row or all(c is None or str(c).strip() == "" for c in row):
                                continue
                            cleaned_row = []
                            for c in row:
                                val = str(c).strip().replace("\n", " ") if c is not None else ""
                                for old_txt, new_txt in replace_map.items():
                                    if old_txt in val:
                                        count = val.count(old_txt)
                                        val = val.replace(old_txt, new_txt)
                                        replacement_count += count
                                cleaned_row.append(val)
                            cleaned_rows.append(cleaned_row)
                        if cleaned_rows and len(cleaned_rows) >= 2:
                            extracted_tables.append({
                                "page": page_idx,
                                "table_idx": t_idx + 1,
                                "rows": cleaned_rows
                            })
    except Exception as e:
        logger.error(f"Failed extracting PDF for edit: {e}")
        return f"Error reading PDF content for edit: {str(e)}"

    # Determine destination
    safe_title = f"{path.stem}_edited"
    if output_path and str(output_path).strip():
        out_dest = _resolve_word_path(output_path, title=safe_title)
    else:
        out_dest = Path.home() / "Desktop" / f"{safe_title}.docx"
        if not out_dest.parent.exists():
            out_dest = DEFAULT_OUTPUT_DIR / f"{safe_title}.docx"

    out_dest.parent.mkdir(parents=True, exist_ok=True)

    # Build updated deliverable
    try:
        doc = Document()
        for section in doc.sections:
            section.top_margin = Cm(2.5)
            section.bottom_margin = Cm(2.5)
            section.left_margin = Cm(3)
            section.right_margin = Cm(2.5)

        # Title
        t = doc.add_heading(f"Document: {path.stem}", level=0)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if t.runs:
            t.runs[0].font.color.rgb = RGBColor(0x1A, 0x56, 0xDB)
            t.runs[0].font.size = Pt(22)

        # Subtitle
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = sub.add_run(f"Updated from PDF '{path.name}' by DeskPilot Agent  •  {datetime.now().strftime('%B %d, %Y')}")
        sub_run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        sub_run.font.size = Pt(9.5)
        sub_run.font.italic = True

        doc.add_paragraph()

        # Body paragraphs from PDF
        for sec in extracted_sections:
            content = sec["content"]
            lines = content.splitlines()
            for line in lines:
                line_str = line.strip()
                if not line_str:
                    continue
                if len(line_str) < 60 and (line_str.isupper() or line_str.startswith("#")):
                    clean_h = line_str.lstrip("# ").strip()
                    h = doc.add_heading(clean_h, level=1)
                    if h.runs:
                        h.runs[0].font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
                        h.runs[0].font.size = Pt(14)
                elif line_str.startswith(("-", "•", "*")):
                    doc.add_paragraph(line_str.lstrip("-•* ").strip(), style="List Bullet")
                else:
                    doc.add_paragraph(line_str)

        # Extracted and modified tables
        for tbl_info in extracted_tables:
            doc.add_heading(f"Table Data (Page {tbl_info['page']})", level=2)
            tbl = doc.add_table(rows=1, cols=len(tbl_info["rows"][0]))
            _apply_table_styling(tbl, tbl_info["rows"])
            doc.add_paragraph()

        # Append additional content if provided
        if append_content and str(append_content).strip():
            append_sections = _parse_input_to_sections(str(append_content).strip())
            for sec in append_sections:
                heading = sec.get("heading", "").strip()
                body = sec.get("content", "").strip()
                table_data = sec.get("table", None)

                if heading:
                    h = doc.add_heading(heading, level=1)
                    if h.runs:
                        h.runs[0].font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
                        h.runs[0].font.size = Pt(14)

                if body:
                    for para_text in body.split("\n\n"):
                        para_text = para_text.strip()
                        if para_text:
                            if para_text.startswith(("-", "•", "*")):
                                for item in para_text.splitlines():
                                    item = item.lstrip("-•* ").strip()
                                    if item:
                                        doc.add_paragraph(item, style="List Bullet")
                            else:
                                doc.add_paragraph(para_text)

                if table_data and isinstance(table_data, list) and len(table_data) >= 2:
                    tbl = doc.add_table(rows=1, cols=len(table_data[0]))
                    _apply_table_styling(tbl, table_data)
                    doc.add_paragraph()

        # Footer
        footer = doc.sections[0].footer
        footer_para = footer.paragraphs[0]
        footer_para.text = f"DeskPilot Agent  •  Edited from {path.name}  •  Deliverable"
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.save(str(out_dest))
        abs_out = str(out_dest.resolve())
        verify_report = verify_word_document(abs_out)

        changes_summary = []
        if replacement_count:
            changes_summary.append(f"Applied {replacement_count} replacement(s)")
        if append_content:
            changes_summary.append("Appended requested content/sections")
        if not changes_summary:
            changes_summary.append("Converted and re-synthesized content")

        resp = [
            f"SUCCESS: PDF '{path.name}' edited and deliverable saved to '{abs_out}'",
            f"Modifications: {'; '.join(changes_summary)}",
        ]
        if instructions:
            resp.append(f"Instructions: {instructions}")
        resp.append(f"\nVerification:\n{verify_report}")
        return "\n".join(resp)

    except Exception as e:
        logger.error(f"Error compiling edited PDF deliverable: {e}")
        return f"Error creating deliverable from PDF '{path.name}': {str(e)}"

