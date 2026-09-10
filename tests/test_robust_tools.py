"""
DeskPilot — Automated Test Suite for Robust Tools & Output Formatting
Validates Excel parsing, formula generation, file read/write tools, and save preferences.
"""

import unittest
import openpyxl
from pathlib import Path
import json
import shutil
import tempfile

from backend.tools.excel_tools import (
    _clean_html_entities,
    _parse_cell_value,
    _parse_headers,
    _parse_table_data,
    create_excel_workbook,
)
from backend.tools.file_tools import read_file, write_file
from backend.utils.save_preferences import (
    get_save_preferences,
    set_save_preferences,
    resolve_effective_save_dir,
)


class TestRobustTools(unittest.TestCase):

    def setUp(self):
        self.temp_dir = Path(tempfile.mkdtemp())

    def tearDown(self):
        if self.temp_dir.exists():
            shutil.rmtree(str(self.temp_dir), ignore_errors=True)

    def test_clean_html_entities(self):
        mangled = "&#91;&#91;'Income', '5000'&#93;, &#91;'Rent', '1500'&#93;&#93;"
        cleaned = _clean_html_entities(mangled)
        self.assertEqual(cleaned, "[['Income', '5000'], ['Rent', '1500']]")

    def test_parse_cell_value(self):
        # Numeric integers
        self.assertEqual(_parse_cell_value("5000"), 5000)
        # Currency formatting
        self.assertEqual(_parse_cell_value("$5,000"), 5000)
        self.assertEqual(_parse_cell_value("$1,250.50"), 1250.50)
        # Percentage
        self.assertEqual(_parse_cell_value("10%"), 0.10)
        # Formulas
        self.assertEqual(_parse_cell_value("=SUM(B2:B5)"), "=SUM(B2:B5)")
        # Booleans
        self.assertEqual(_parse_cell_value("true"), True)
        self.assertEqual(_parse_cell_value("False"), False)
        # Strings
        self.assertEqual(_parse_cell_value("Mortgage / Rent"), "Mortgage / Rent")

    def test_parse_table_data_mangled_html(self):
        # Emulating exactly what Bedrock Nova Pro sent in user Screenshot 1
        mangled_input = (
            "&#91;&#91;'Income', 'Salary', '5000'&#93;, "
            "&#91;'Income', 'Bonus', '1000'&#93;,, "
            "&#91;'Expenses', 'Rent', '1500'&#93;&#93;"
        )
        rows = _parse_table_data(mangled_input, expected_cols=3)
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[0], ["Income", "Salary", 5000])
        self.assertEqual(rows[1], ["Income", "Bonus", 1000])
        self.assertEqual(rows[2], ["Expenses", "Rent", 1500])

    def test_parse_table_data_python_single_quotes(self):
        py_literal = "[['Salary', 5000], ['Housing', 1600], ['Groceries', 650]]"
        rows = _parse_table_data(py_literal, expected_cols=2)
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[0], ["Salary", 5000])
        self.assertEqual(rows[1], ["Housing", 1600])
        self.assertEqual(rows[2], ["Groceries", 650])

    def test_parse_table_data_flat_list_chunking(self):
        # Flat 1D list when 2 columns expected
        flat_list = ["Rent", 1500, "Food", 500, "Utilities", 250]
        rows = _parse_table_data(flat_list, expected_cols=2)
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[0], ["Rent", 1500])
        self.assertEqual(rows[1], ["Food", 500])
        self.assertEqual(rows[2], ["Utilities", 250])

    def test_create_excel_workbook_end_to_end(self):
        target_path = self.temp_dir / "test_budget.xlsx"
        res = create_excel_workbook(
            sheet_name="Test Budget",
            headers=["Category", "Amount ($)"],
            rows="&#91;&#91;'Salary', '5000'&#93;, &#91;'Rent', '1500'&#93;, &#91;'Groceries', '600'&#93;&#93;",
            output_path=str(target_path),
            title="Monthly Family Budget",
            add_totals=True,
        )

        self.assertTrue("SUCCESS" in res)
        self.assertTrue(target_path.exists())

        # Inspect workbook with openpyxl
        wb = openpyxl.load_workbook(str(target_path))
        ws = wb.active
        self.assertEqual(ws.title, "Test Budget")

        # Row 1 is Title (Merged)
        self.assertEqual(ws.cell(row=1, column=1).value, "Monthly Family Budget")

        # Row 3 is Headers
        self.assertEqual(ws.cell(row=3, column=1).value, "Category")
        self.assertEqual(ws.cell(row=3, column=2).value, "Amount ($)")

        # Row 4 is first data row
        self.assertEqual(ws.cell(row=4, column=1).value, "Salary")
        self.assertEqual(ws.cell(row=4, column=2).value, 5000)
        self.assertIsInstance(ws.cell(row=4, column=2).value, int)

        # Row 5 is second data row
        self.assertEqual(ws.cell(row=5, column=1).value, "Rent")
        self.assertEqual(ws.cell(row=5, column=2).value, 1500)

        # Row 6 is third data row
        self.assertEqual(ws.cell(row=6, column=1).value, "Groceries")
        self.assertEqual(ws.cell(row=6, column=2).value, 600)

        # Row 7 is TOTAL row
        self.assertEqual(ws.cell(row=7, column=1).value, "TOTAL")
        self.assertEqual(ws.cell(row=7, column=2).value, "=SUM(B4:B6)")

    def test_read_and_write_file_tools(self):
        target_file = self.temp_dir / "test_notes.md"
        content = "# Project Notes\n\n- Task 1 verified\n- Task 2 verified\n"

        # 1. Write file
        res_write = write_file(str(target_file), content)
        self.assertTrue("SUCCESS" in res_write)
        self.assertTrue(target_file.exists())

        # 2. Read file
        read_content = read_file(str(target_file))
        self.assertEqual(read_content, content)

    def test_save_preferences(self):
        # Set preference
        set_save_preferences("Documents", always_save=True)
        prefs = get_save_preferences()
        self.assertEqual(prefs["save_location"], "Documents")
        self.assertTrue(prefs["always_save"])

        # Reset to Desktop
        set_save_preferences("Desktop", always_save=True)
        prefs = get_save_preferences()
        self.assertEqual(prefs["save_location"], "Desktop")


if __name__ == "__main__":
    unittest.main()
