"""
DeskPilot — Automated Unit Tests for Document Intelligence Suite
Tests reading and editing capabilities for Word (.docx), PDF (.pdf), and Excel (.xlsx) files.
"""

import unittest
import tempfile
import shutil
import json
from pathlib import Path
from docx import Document
import openpyxl

from backend.utils.document_resolver import resolve_document_path
from backend.tools.word_tools import read_word_document, edit_word_document, create_word_report, verify_word_document
from backend.tools.pdf_tools import read_pdf, edit_pdf_document, list_pdfs_in_folder
from backend.tools.excel_tools import read_excel_file, edit_excel_file, create_excel_workbook, verify_excel_workbook
from backend.tools import get_tools_for_agent, TOOL_REGISTRY
from backend.utils.trust import classify_action, TrustTier


class TestDocumentEditingSuite(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.temp_path = Path(self.temp_dir)

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_document_resolver_with_samples_and_stems(self):
        """Test resolving files by full path, bare filename, and stem without extension."""
        # Test finding sample invoice PDF
        p1 = resolve_document_path("invoice_001_acme.pdf")
        self.assertIsNotNone(p1)
        self.assertTrue(p1.exists())
        self.assertEqual(p1.suffix.lower(), ".pdf")

        # Test finding sample supplier ledger Excel without extension
        p2 = resolve_document_path("supplier_ledger", default_ext=".xlsx")
        self.assertIsNotNone(p2)
        self.assertTrue(p2.exists())
        self.assertEqual(p2.suffix.lower(), ".xlsx")

        # Test non-existent file
        p3 = resolve_document_path("non_existent_file_xyz_123.docx")
        self.assertIsNone(p3)

    def test_word_read_and_edit_workflow(self):
        """Test reading an existing Word document and applying text replacements and appended sections."""
        # 1. Create initial Word document
        sample_docx = self.temp_path / "Medical_Lab_Report.docx"
        doc = Document()
        doc.add_heading("Patient Medical Assessment", level=0)
        p1 = doc.add_paragraph("Patient Status: Preliminary Positive for TB")
        p2 = doc.add_paragraph("Notes: Further microscopic analysis pending.")
        
        # Add a table
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Test Name"
        table.rows[0].cells[1].text = "Result"
        table.rows[1].cells[0].text = "Acid-Fast Stain"
        table.rows[1].cells[1].text = "Pending"
        doc.save(str(sample_docx))

        # 2. Read Word Document
        read_res = read_word_document(str(sample_docx))
        self.assertIn("Patient Medical Assessment", read_res)
        self.assertIn("Preliminary Positive for TB", read_res)
        self.assertIn("Acid-Fast Stain", read_res)

        # 3. Edit Word Document (text replacement in paragraph and table, plus append section)
        replacements = {
            "Preliminary Positive for TB": "CLEARED - Negative for TB",
            "Pending": "Confirmed Negative"
        }
        append_md = (
            "## Physician Signoff\n"
            "- Final validation completed.\n"
            "- Discharged with full clearance.\n\n"
            "| Metric | Status |\n"
            "| --- | --- |\n"
            "| Discharge | Approved |\n"
        )
        
        edit_res = edit_word_document(
            file_path=str(sample_docx),
            instructions="Update TB test results to Negative and append physician signoff",
            replacements=replacements,
            append_content=append_md
        )

        self.assertIn("SUCCESS", edit_res)
        self.assertIn("Safety Backup Created", edit_res)

        # Verify edited document
        updated_doc = Document(str(sample_docx))
        all_text = " ".join(p.text for p in updated_doc.paragraphs)
        self.assertIn("CLEARED - Negative for TB", all_text)
        self.assertNotIn("Preliminary Positive for TB", all_text)
        self.assertIn("Final validation completed", all_text)

        # Check table edit
        table_cells = [c.text for r in updated_doc.tables[0].rows for c in r.cells]
        self.assertIn("Confirmed Negative", table_cells)

        # Check safety backup exists
        backups = list(self.temp_path.glob("*.bak"))
        self.assertTrue(len(backups) >= 1)

    def test_word_edit_to_custom_output_path(self):
        """Test editing a Word doc and saving to a custom destination file."""
        src_docx = self.temp_path / "Original_Notes.docx"
        doc = Document()
        doc.add_heading("Original Notes", level=1)
        doc.add_paragraph("Draft version 1.0")
        doc.save(str(src_docx))

        out_docx = self.temp_path / "Final_Notes.docx"
        res = edit_word_document(
            file_path=str(src_docx),
            replacements={"Draft version 1.0": "Production Release 2.0"},
            output_path=str(out_docx)
        )

        self.assertIn("SUCCESS", res)
        self.assertTrue(out_docx.exists())
        
        # Check that original remains intact and new file has replacement
        doc_orig = Document(str(src_docx))
        self.assertIn("Draft version 1.0", doc_orig.paragraphs[1].text)

        doc_out = Document(str(out_docx))
        self.assertIn("Production Release 2.0", doc_out.paragraphs[1].text)

    def test_excel_read_and_edit_workflow(self):
        """Test reading an existing Excel file and updating cells and appending rows."""
        # 1. Create initial Excel workbook
        sample_xlsx = self.temp_path / "quarterly_budget.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Financials"
        ws.append(["Department", "Allocated", "Status"])
        ws.append(["Engineering", 50000, "Draft"])
        ws.append(["Marketing", 20000, "Draft"])
        wb.save(str(sample_xlsx))

        # 2. Read Excel File
        read_res = read_excel_file(str(sample_xlsx))
        data = json.loads(read_res)
        self.assertEqual(data["sheet"], "Financials")
        self.assertEqual(len(data["rows"]), 2)

        # 3. Edit Excel File: Update Marketing to 25000 and status Approved, and append row for Operations
        cell_updates = [
            {"cell": "B3", "value": 25000},
            {"cell": "C3", "value": "Approved"}
        ]
        append_rows = [
            ["Operations", 15000, "Approved"]
        ]

        edit_res = edit_excel_file(
            file_path=str(sample_xlsx),
            cell_updates=cell_updates,
            append_rows=append_rows
        )

        self.assertIn("SUCCESS", edit_res)
        self.assertIn("Safety Backup Created", edit_res)

        # Verify changes in workbook
        updated_wb = openpyxl.load_workbook(str(sample_xlsx))
        updated_ws = updated_wb["Financials"]
        self.assertEqual(updated_ws["B3"].value, 25000)
        self.assertEqual(updated_ws["C3"].value, "Approved")
        self.assertEqual(updated_ws["A4"].value, "Operations")
        self.assertEqual(updated_ws["B4"].value, 15000)

        # Verify backup was created
        backups = list(self.temp_path.glob("*.bak"))
        self.assertTrue(len(backups) >= 1)

    def test_pdf_edit_and_deliverable_generation(self):
        """Test editing a sample PDF and compiling an updated deliverable document."""
        sample_pdf = Path("sample_data/invoice_001_acme.pdf")
        if not sample_pdf.exists():
            self.skipTest("sample_data/invoice_001_acme.pdf not found")

        out_docx = self.temp_path / "Acme_Invoice_Updated.docx"
        res = edit_pdf_document(
            file_path=str(sample_pdf),
            instructions="Update supplier discount and add approval note",
            replacements={"Acme Office Solutions Ltd.": "Acme Global Industries"},
            append_content="## Accounts Payable Approval\n- Payment authorized on Net-30 terms.",
            output_path=str(out_docx)
        )

        self.assertIn("SUCCESS", res)
        self.assertTrue(out_docx.exists())

        # Inspect compiled deliverable Word doc
        doc = Document(str(out_docx))
        doc_text = " ".join(p.text for p in doc.paragraphs)
        self.assertIn("Acme Global Industries", doc_text)
        self.assertIn("Payment authorized on Net-30 terms", doc_text)

    def test_trust_tier_classification(self):
        """Verify trust tiers for read vs edit actions."""
        # Reads are Green
        self.assertEqual(classify_action("read_word_document"), TrustTier.GREEN)
        self.assertEqual(classify_action("read_pdf"), TrustTier.GREEN)
        self.assertEqual(classify_action("read_excel_file"), TrustTier.GREEN)

        # Edits are Yellow
        self.assertEqual(classify_action("edit_word_document"), TrustTier.YELLOW)
        self.assertEqual(classify_action("edit_pdf_document"), TrustTier.YELLOW)
        self.assertEqual(classify_action("edit_excel_file"), TrustTier.YELLOW)

    def test_agent_tool_resolution(self):
        """Verify that agents resolve new document tools without errors."""
        personal_tools = get_tools_for_agent([
            "read_word_document", "edit_word_document",
            "read_pdf", "edit_pdf_document",
            "read_excel_file", "edit_excel_file"
        ])
        self.assertEqual(len(personal_tools), 6)


if __name__ == "__main__":
    unittest.main()
